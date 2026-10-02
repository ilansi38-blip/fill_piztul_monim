import argparse
import json
import re
from pathlib import Path

import docx
from docx.oxml.ns import qn

TEMPLATES_DIR = "templates"
CONFIGS_DIR = "configs"

PLACEHOLDER_PATTERN = re.compile(r"__(.+?)__")


def _own_text_nodes(paragraph_xml):
    """w:t descendants of paragraph_xml, excluding ones that actually belong to
    a nested paragraph (e.g. a text box's w:txbxContent can sit inside a run
    of an outer paragraph -- its w:t must stay with the inner paragraph only,
    or the outer paragraph's text gets the inner one's content glued onto it)."""
    nodes = []
    for t in paragraph_xml.findall(".//" + qn("w:t")):
        ancestor = t.getparent()
        while ancestor is not None and ancestor.tag != qn("w:p"):
            ancestor = ancestor.getparent()
        if ancestor is paragraph_xml:
            nodes.append(t)
    return nodes


def _iter_paragraphs(document):
    """Yield (text_nodes, full_text) for every paragraph in the document.

    Walking document.element.body.iter(w:p) instead of document.paragraphs /
    document.tables also reaches paragraphs nested inside text boxes
    (w:drawing/.../w:txbxContent), which the high-level python-docx API skips.
    """
    for paragraph_xml in document.element.body.iter(qn("w:p")):
        text_nodes = _own_text_nodes(paragraph_xml)
        full_text = "".join(node.text or "" for node in text_nodes)
        yield text_nodes, full_text


def extract_keywords(docx_path: str) -> dict:
    """Return {keyword: ""} for every __keyword__ placeholder found in docx_path."""
    document = docx.Document(docx_path)
    keywords = {}
    for _, text in _iter_paragraphs(document):
        for key in PLACEHOLDER_PATTERN.findall(text):
            keywords[key.strip()] = ""
    return keywords


def load_config(config_path: str) -> dict:
    try:
        with open(config_path, "r", encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        return {}


def save_config(config: dict, config_path: str) -> None:
    Path(config_path).parent.mkdir(parents=True, exist_ok=True)
    with open(config_path, "w", encoding="utf-8") as f:
        json.dump(config, f, ensure_ascii=False, indent=2)


def merge_keywords(docx_path: str, config_path: str) -> dict:
    """Parse docx_path and add any newly discovered keys to the config file,
    keeping whatever value an already-known key currently has."""
    found = extract_keywords(docx_path)
    config = load_config(config_path)
    for key, default_value in found.items():
        config.setdefault(key, default_value)
    save_config(config, config_path)
    return config


def config_path_for(template_path, configs_dir: str = CONFIGS_DIR) -> str:
    """The configs/<template stem>.json path that goes with a given template."""
    return str(Path(configs_dir) / f"{Path(template_path).stem}.json")


def discover_templates(templates_dir: str = TEMPLATES_DIR) -> list:
    """All .docx templates in templates_dir (skips Word's ~$ lock files)."""
    return sorted(
        p for p in Path(templates_dir).glob("*.docx") if not p.name.startswith("~$")
    )


def sync_templates(templates_dir: str = TEMPLATES_DIR, configs_dir: str = CONFIGS_DIR) -> dict:
    """Merge every template's placeholders into its own configs/<stem>.json.
    Adding a new template file is all that's needed -- run this (or `python
    parser.py`) afterwards and its config file appears/updates automatically.
    Returns {template_stem: config}.
    """
    results = {}
    for template in discover_templates(templates_dir):
        config_path = config_path_for(template, configs_dir)
        results[template.stem] = merge_keywords(str(template), config_path)
    return results


def _nearest_ancestor(node, tag):
    ancestor = node.getparent()
    while ancestor is not None and ancestor.tag != tag:
        ancestor = ancestor.getparent()
    return ancestor


def fill_document(docx_path: str, values: dict, output_path: str) -> str:
    """Fill __keyword__ placeholders in docx_path using values, saving as a
    separate copy at output_path (docx_path itself is left untouched).
    Placeholders with no non-empty value in `values` are stripped out (left
    blank) rather than left as literal "__key__" text. Table rows where every
    placeholder in the row ended up with no value are deleted entirely (rows
    with no placeholders at all, e.g. header/checkbox rows, are never touched).

    A placeholder can be split across multiple w:t runs with different
    formatting (e.g. a leading layout/indent run before the label vs. the run
    that actually holds "__key__"). Only the runs the placeholder's own
    characters overlap are touched -- the replacement text goes into whichever
    run holds the placeholder's start, so it keeps that run's formatting
    instead of picking up an unrelated run's font/size.
    """
    document = docx.Document(docx_path)

    row_stats = {}  # w:tr element -> {"has_placeholder": bool, "has_value": bool}

    for text_nodes, full_text in _iter_paragraphs(document):
        if "__" not in full_text:
            continue

        matches = list(PLACEHOLDER_PATTERN.finditer(full_text))
        if not matches:
            continue

        row = _nearest_ancestor(text_nodes[0], qn("w:tr")) if text_nodes else None

        texts = [node.text or "" for node in text_nodes]
        offsets = []
        pos = 0
        for t in texts:
            offsets.append(pos)
            pos += len(t)

        new_texts = list(texts)
        for match in matches:
            key = match.group(1).strip()
            value = values.get(key)
            replacement = str(value) if value else ""

            if row is not None:
                stats = row_stats.setdefault(row, {"has_placeholder": False, "has_value": False})
                stats["has_placeholder"] = True
                if value:
                    stats["has_value"] = True

            m_start, m_end = match.start(), match.end()
            inserted = False
            for i, t in enumerate(texts):
                n_start = offsets[i]
                n_end = n_start + len(t)
                if n_end <= m_start or n_start >= m_end:
                    continue  # this run doesn't overlap the placeholder at all
                prefix = t[: m_start - n_start] if n_start < m_start else ""
                suffix = t[m_end - n_start:] if n_end > m_end else ""
                new_texts[i] = prefix + (replacement if not inserted else "") + suffix
                inserted = True

        for node, new_t in zip(text_nodes, new_texts):
            node.text = new_t

    for row, stats in row_stats.items():
        if stats["has_placeholder"] and not stats["has_value"]:
            parent = row.getparent()
            if parent is not None:
                parent.remove(row)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def main():
    arg_parser = argparse.ArgumentParser(
        description="Scan .docx template(s) for __keyword__ placeholders "
        "(including ones inside text boxes) and merge them into a matching "
        "configs/<template>.json file. With no path given, syncs every "
        "template found in templates/."
    )
    arg_parser.add_argument(
        "docx_path", nargs="?", default=None,
        help="A single template to sync. Omit to sync all of templates/.",
    )
    arg_parser.add_argument("--templates-dir", default=TEMPLATES_DIR)
    arg_parser.add_argument("--configs-dir", default=CONFIGS_DIR)
    args = arg_parser.parse_args()

    if args.docx_path:
        config_path = config_path_for(args.docx_path, args.configs_dir)
        config = merge_keywords(args.docx_path, config_path)
        results = {Path(args.docx_path).stem: config}
    else:
        results = sync_templates(args.templates_dir, args.configs_dir)

    for stem, config in results.items():
        print(f"{stem} -> {len(config)} keys ({config_path_for(stem + '.docx', args.configs_dir)}):")
        for key, value in config.items():
            print(f"  {key}: {value!r}")


if __name__ == "__main__":
    main()
